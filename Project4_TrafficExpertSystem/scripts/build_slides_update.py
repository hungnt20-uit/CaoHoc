"""Sinh bộ slide cập nhật (PPTX) khớp source hiện tại — NĐ 168/2024.

Chạy:  python3 scripts/build_slides_update.py
Output: docs/BaoCao_TrafficES_Slides_Update.pptx
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "BaoCao_TrafficES_Slides_Update.pptx"

# Palette — UIT / báo cáo học thuật (tránh purple-on-white)
NAVY = RGBColor(0x0B, 0x25, 0x40)
TEAL = RGBColor(0x1B, 0x6B, 0x7A)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)
LIGHT = RGBColor(0xF4, 0xF7, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x4A, 0x55, 0x68)
DARK = RGBColor(0x1A, 0x1A, 0x2E)
GREEN = RGBColor(0x1F, 0x7A, 0x4D)
RED = RGBColor(0xA3, 0x2D, 0x2D)

W, H = Inches(13.333), Inches(7.5)  # 16:9


def _set_run(run, text, size=18, bold=False, color=DARK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        from lxml import etree

        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", "Arial")


def _add_textbox(slide, left, top, width, height, text, *, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    _set_run(p.add_run() if not p.runs else p.runs[0], text, size=size, bold=bold, color=color)
    if not p.runs:
        run = p.add_run()
        _set_run(run, text, size=size, bold=bold, color=color)
    else:
        # first run may be empty from add_textbox
        if p.runs[0].text == "" and len(p.runs) == 1:
            _set_run(p.runs[0], text, size=size, bold=bold, color=color)
        elif p.runs[0].text != text:
            _set_run(p.runs[0], text, size=size, bold=bold, color=color)
    return box


def add_text(slide, left, top, width, height, lines, *, size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT, space=1.05):
    """lines: str | list[str | tuple(text, bold)]"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if isinstance(lines, str):
        lines = [lines]
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, is_bold = item
        else:
            text, is_bold = item, bold
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4 * space)
        run = p.add_run()
        _set_run(run, text, size=size, bold=is_bold, color=color)
    return box


def fill_shape(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def bar(slide, top=None):
    """Top navy bar + teal accent line."""
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, W, Inches(0.85))
    fill_shape(shp, NAVY)
    accent = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, Inches(0.85), W, Inches(0.06)
    )
    fill_shape(accent, TEAL)


def footer(slide, page: int, total: int = 22):
    add_text(
        slide,
        Inches(0.5),
        Inches(7.1),
        Inches(10),
        Inches(0.3),
        "CS2307.CH201 · Nhóm 4 · Traffic Expert System · NĐ 168/2024",
        size=10,
        color=MUTED,
    )
    add_text(
        slide,
        Inches(11.8),
        Inches(7.1),
        Inches(1.2),
        Inches(0.3),
        str(page),
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def title_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, W, H)
    fill_shape(bg, NAVY)
    accent = s.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, Inches(5.9), W, Inches(0.12)
    )
    fill_shape(accent, TEAL)
    add_text(
        s,
        Inches(0.8),
        Inches(1.2),
        Inches(11.5),
        Inches(0.5),
        "ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH · TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN",
        size=14,
        color=RGBColor(0xA8, 0xC5, 0xD4),
        align=PP_ALIGN.CENTER,
    )
    add_text(
        s,
        Inches(0.8),
        Inches(1.9),
        Inches(11.5),
        Inches(0.4),
        "BÁO CÁO CUỐI KỲ — CẬP NHẬT",
        size=16,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        s,
        Inches(0.8),
        Inches(2.5),
        Inches(11.5),
        Inches(1.4),
        "HỆ THỐNG CHUYÊN GIA TƯ VẤN HÀNH VI\nVI PHẠM LUẬT GIAO THÔNG\nDỰA TRÊN REASONING",
        size=28,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        s,
        Inches(0.8),
        Inches(4.3),
        Inches(11.5),
        Inches(0.8),
        [
            "Chuyên đề: Nghiên cứu và Ứng dụng về Công nghệ Tri thức",
            "GVHD: TS. Nguyễn Đình Hiển  ·  Lớp: CS2307.CH201  ·  Nhóm: 4",
            "KB: Nghị định 168/2024/NĐ-CP  ·  Kiến trúc: Neuro-Symbolic (LLM + Engine ký hiệu)",
        ],
        size=14,
        color=RGBColor(0xC8, 0xD8, 0xE0),
        align=PP_ALIGN.CENTER,
    )
    add_text(
        s,
        Inches(0.8),
        Inches(6.3),
        Inches(11.5),
        Inches(0.4),
        "TP. Hồ Chí Minh, Tháng 08/2026",
        size=13,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def members_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(10), Inches(0.5), "THÀNH VIÊN NHÓM 4", size=26, bold=True, color=WHITE)
    members = [
        ("1", "250101025", "Nguyễn Thanh Hùng"),
        ("2", "250101065", "Bùi Phạm Minh Thi"),
        ("3", "250101055", "Võ Anh Quân"),
        ("4", "250101039", "Trần Hoàng Minh"),
        ("5", "250101049", "Trần Thảo Nguyên"),
    ]
    headers = ["STT", "Mã học viên", "Họ và tên"]
    # header row
    y0 = Inches(1.4)
    cols = [Inches(1.2), Inches(3.5), Inches(5.5)]
    x0 = Inches(1.5)
    for i, h in enumerate(headers):
        cell = s.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, x0 + sum(cols[:i], Inches(0)), y0, cols[i], Inches(0.55)
        )
        fill_shape(cell, TEAL)
        add_text(s, x0 + sum(cols[:i], Inches(0)), y0 + Inches(0.1), cols[i], Inches(0.4), h, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    for r, (stt, ma, ten) in enumerate(members):
        y = y0 + Inches(0.55) * (r + 1)
        bgc = LIGHT if r % 2 == 0 else WHITE
        for i, val in enumerate((stt, ma, ten)):
            cell = s.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RECTANGLE, x0 + sum(cols[:i], Inches(0)), y, cols[i], Inches(0.55)
            )
            fill_shape(cell, bgc)
            cell.line.color.rgb = RGBColor(0xD0, 0xD8, 0xDE)
            add_text(
                s,
                x0 + sum(cols[:i], Inches(0)),
                y + Inches(0.1),
                cols[i],
                Inches(0.4),
                val,
                size=15,
                color=DARK,
                align=PP_ALIGN.CENTER,
            )
    footer(s, 2)


def overview_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "TỔNG QUAN DỰ ÁN", size=26, bold=True, color=WHITE)
    cards = [
        ("VẤN ĐỀ", "Xử lý tự động mô tả vi phạm giao thông tiếng Việt phức tạp — cần suy luận ngữ nghĩa, minh bạch pháp lý và chính xác cao."),
        ("THÁCH THỨC", "Luật nhiều tầng (NĐ 168/2024), ngôn ngữ khẩu ngữ đa dạng, không được ảo giác căn cứ / mức phạt."),
        ("MỤC ĐÍCH", "Hệ chuyên gia hybrid: LLM/NLU trích facts + engine ký hiệu suy diễn mức phạt, căn cứ và chuỗi giải thích."),
        ("CẬP NHẬT", "KB chuyển sang NĐ 168/2024; UI Streamlit + Nodes graph Legal-Onto; hỗ trợ OpenAI / Claude / Local LLM."),
    ]
    positions = [(0.5, 1.2), (6.9, 1.2), (0.5, 4.0), (6.9, 4.0)]
    for (title, body), (x, y) in zip(cards, positions):
        card = s.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.9), Inches(2.5)
        )
        fill_shape(card, LIGHT)
        card.line.color.rgb = TEAL
        add_text(s, Inches(x + 0.3), Inches(y + 0.25), Inches(5.3), Inches(0.4), title, size=16, bold=True, color=TEAL)
        add_text(s, Inches(x + 0.3), Inches(y + 0.8), Inches(5.3), Inches(1.5), body, size=14, color=DARK)
    footer(s, 3)


def goals_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "MỤC TIÊU HỆ THỐNG", size=26, bold=True, color=WHITE)
    goals = [
        ("01", "Nhận diện tình huống", "Nhận mô tả NL tiếng Việt → trích facts, suy tình tiết tăng/giảm nặng (heuristic hoặc LLM)."),
        ("02", "Quyết định xử phạt", "Engine ký hiệu áp rules YAML (NĐ 168) → tiền phạt, tước GPLX, căn cứ Điều/Khoản/điểm."),
        ("03", "Giải thích minh bạch", "Trace đầy đủ + Legal-Onto validator: LLM không được sinh số liệu / căn cứ pháp lý."),
    ]
    for i, (num, title, body) in enumerate(goals):
        x = Inches(0.5 + i * 4.2)
        card = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(3.9), Inches(4.6))
        fill_shape(card, WHITE)
        card.line.color.rgb = RGBColor(0xD0, 0xD8, 0xDE)
        circle = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x + Inches(1.35), Inches(1.9), Inches(1.2), Inches(1.2))
        fill_shape(circle, TEAL)
        add_text(s, x + Inches(1.35), Inches(2.2), Inches(1.2), Inches(0.7), num, size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.25), Inches(3.4), Inches(3.4), Inches(0.5), title, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.3), Inches(4.1), Inches(3.3), Inches(1.7), body, size=13, color=DARK, align=PP_ALIGN.CENTER)
    footer(s, 4)


def principles_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "NGUYÊN TẮC KIẾN TRÚC", size=26, bold=True, color=WHITE)
    left = [
        ("Chống ảo giác pháp lý", "Validator / grounding + engine ký hiệu phê duyệt mọi kết luận. LLM không sinh tiền phạt hay Điều luật."),
        ("Ánh xạ thuật ngữ", "Keyphrases đời thường → concept Legal-Onto → fact trong Working Memory."),
    ]
    right = [
        ("Quyết định thuộc Engine", "Forward chaining + A* + meta-rule. Symbolic Engine đưa ra mọi quyết định xử phạt."),
        ("Truy xuất 100%", "Mỗi kết luận gắn can_cu (NĐ, Điều, Khoản, điểm) và TraceStep (RULE / META / KL)."),
    ]
    for i, (title, body) in enumerate(left):
        y = Inches(1.3 + i * 2.5)
        card = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.5), y, Inches(5.9), Inches(2.2))
        fill_shape(card, LIGHT)
        add_text(s, Inches(0.8), y + Inches(0.3), Inches(5.3), Inches(0.4), title, size=16, bold=True, color=ACCENT)
        add_text(s, Inches(0.8), y + Inches(0.85), Inches(5.3), Inches(1.1), body, size=14, color=DARK)
    for i, (title, body) in enumerate(right):
        y = Inches(1.3 + i * 2.5)
        card = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.9), y, Inches(5.9), Inches(2.2))
        fill_shape(card, LIGHT)
        add_text(s, Inches(7.2), y + Inches(0.3), Inches(5.3), Inches(0.4), title, size=16, bold=True, color=TEAL)
        add_text(s, Inches(7.2), y + Inches(0.85), Inches(5.3), Inches(1.1), body, size=14, color=DARK)
    footer(s, 5)


def arch_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "KIẾN TRÚC NEURO-SYMBOLIC (KHỚP SOURCE)", size=24, bold=True, color=WHITE)
    layers = [
        ("UI", "Streamlit\n+ Nodes graph", NAVY),
        ("Service", "TrafficESService\nanswer / continue", TEAL),
        ("NLU", "Extract · Infer\nKG · Ground · Clarify", ACCENT),
        ("Engine", "A* · FC Rules\nMeta · Aggregate", GREEN),
        ("KB", "concepts.yaml\nrules/*.yaml", MUTED),
    ]
    for i, (name, desc, col) in enumerate(layers):
        x = Inches(0.4 + i * 2.55)
        box = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, Inches(1.4), Inches(2.35), Inches(2.4))
        fill_shape(box, col)
        add_text(s, x, Inches(1.6), Inches(2.35), Inches(0.5), name, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.1), Inches(2.3), Inches(2.15), Inches(1.3), desc, size=13, color=WHITE, align=PP_ALIGN.CENTER)
        if i < len(layers) - 1:
            add_text(s, x + Inches(2.15), Inches(2.3), Inches(0.4), Inches(0.4), "→", size=20, bold=True, color=NAVY)
    add_text(
        s,
        Inches(0.5),
        Inches(4.2),
        Inches(12.3),
        Inches(2.3),
        [
            "Luồng dữ liệu:",
            "NL Input → Normalize → Extractor (Heuristic | LLM) → Circumstance Inferrer → Legal-Onto subgraph match",
            "→ Grounding → (Clarify nếu thiếu slot) → Working Memory → A* → Forward chaining → Meta → Kết luận + Trace",
            "",
            "Nguyên tắc hybrid: LLM chỉ sinh facts/tình tiết; Engine ký hiệu sinh mức phạt & căn cứ (chống legal hallucination).",
        ],
        size=14,
        color=DARK,
    )
    footer(s, 6)


def legal_onto_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "MÔ HÌNH LEGAL-ONTO", size=26, bold=True, color=WHITE)
    add_text(
        s,
        Inches(0.5),
        Inches(1.2),
        Inches(12.3),
        Inches(0.8),
        "𝒦 = (Conc, Rel, Rules) ⊕ (Keyphrases, Rela)",
        size=22,
        bold=True,
        color=TEAL,
        align=PP_ALIGN.CENTER,
    )
    items = [
        ("Conc", "Khái niệm pháp lý: phương tiện, hành vi, chỉ số, tình tiết… (concepts.yaml)"),
        ("Rel / Rela", "Trong triển khai: HAS_KEYPHRASE, MAPS_TO fact, HAS_GROUP — star graph theo concept"),
        ("Rules", "Luật điều kiện–kết luận YAML → mức xử phạt + can_cu NĐ 168/2024"),
        ("Keyphrases", "Cụm từ đời thường khớp TF-IDF phrase-anchored → concept → fact WM"),
    ]
    for i, (k, v) in enumerate(items):
        y = Inches(2.2 + i * 1.0)
        tag = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.5), y, Inches(2.2), Inches(0.75))
        fill_shape(tag, TEAL)
        add_text(s, Inches(0.5), y + Inches(0.18), Inches(2.2), Inches(0.45), k, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(s, Inches(3.0), y + Inches(0.18), Inches(9.5), Inches(0.55), v, size=14, color=DARK)
    footer(s, 7)


def concept_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "CONCEPT — 5 THÀNH PHẦN", size=26, bold=True, color=WHITE)
    add_text(s, Inches(0.5), Inches(1.15), Inches(6), Inches(0.4), "Định nghĩa", size=16, bold=True, color=TEAL)
    defs = [
        "Name — tên định danh khái niệm",
        "Content — mô tả nội hàm pháp lý",
        "InnerRul — căn cứ / điều kiện nội tại",
        "Attrs — nhom, fact, value → WM",
        "Keyphrases — cụm nhận diện trong NL",
    ]
    add_text(s, Inches(0.5), Inches(1.6), Inches(6), Inches(3), defs, size=15, color=DARK)
    card = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(1.2), Inches(5.8), Inches(5.2))
    fill_shape(card, LIGHT)
    add_text(s, Inches(7.3), Inches(1.4), Inches(5.2), Inches(0.4), "Ví dụ trong source: xe mô tô", size=15, bold=True, color=NAVY)
    add_text(
        s,
        Inches(7.3),
        Inches(2.0),
        Inches(5.2),
        Inches(4.0),
        [
            'name: "xe mô tô"',
            "content: Phương tiện 2–3 bánh…",
            "inner_rul: Điều 3.39 QCVN…",
            'attrs: nhom=phuong_tien',
            "        fact=phuongtien.loai",
            "        value=xe_may",
            'keyphrases: ["xe máy", "mô tô",',
            '             "xe gắn máy", …]',
        ],
        size=13,
        color=DARK,
    )
    footer(s, 8)


def catalog_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "DANH MỤC CONCEPT TRONG KB", size=26, bold=True, color=WHITE)
    rows = [
        ("Phương tiện", "xe ô tô, xe mô tô → phuongtien.loai"),
        ("Chỉ số", "nồng độ cồn, tốc độ (cần số đo / khung mức)"),
        ("Hành vi", "vượt đèn đỏ, không mũ BH, không dây AT, không GPLX"),
        ("Tình tiết", "gây tai nạn, tái phạm (meta tăng nặng)"),
        ("Hình phạt", "tiền + tước GPLX / trừ điểm (từ ket_luan rule)"),
    ]
    for i, (k, v) in enumerate(rows):
        y = Inches(1.25 + i * 1.0)
        left = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.5), y, Inches(3.0), Inches(0.8))
        fill_shape(left, NAVY if i % 2 == 0 else TEAL)
        add_text(s, Inches(0.5), y + Inches(0.2), Inches(3.0), Inches(0.45), k, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        right = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(3.7), y, Inches(9.0), Inches(0.8))
        fill_shape(right, LIGHT)
        add_text(s, Inches(4.0), y + Inches(0.2), Inches(8.5), Inches(0.45), v, size=14, color=DARK)
    footer(s, 9)


def kg_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "RELATION & KNOWLEDGE GRAPH (THỰC TẾ SOURCE)", size=22, bold=True, color=WHITE)
    add_text(
        s,
        Inches(0.5),
        Inches(1.2),
        Inches(12.3),
        Inches(1.2),
        "KB hiện tại là các ngôi sao concept (không phải đồ thị 1600+ đỉnh). UI nối Legal-Onto → nhóm → concept → keyphrase/fact để trực quan hóa liên thông.",
        size=15,
        color=DARK,
    )
    edges = [
        ("HAS_KEYPHRASE", "Concept sở hữu cụm từ — độ tin theo trọng số TF-IDF"),
        ("EXHIBITS", "Concept được câu hỏi kích hoạt (sau khi phân tích)"),
        ("MAPS_TO", "Concept → fact trong Working Memory"),
        ("HAS_GROUP / HAS_CONCEPT", "Hub ontology gom nhóm (UI nodes graph)"),
        ("MATCHES → FIRES → CONCLUDES", "Chuỗi suy diễn theo tình huống trên UI"),
    ]
    for i, (e, d) in enumerate(edges):
        y = Inches(2.5 + i * 0.75)
        add_text(s, Inches(0.6), y, Inches(4.5), Inches(0.55), e, size=14, bold=True, color=TEAL)
        add_text(s, Inches(5.2), y, Inches(7.5), Inches(0.55), d, size=14, color=DARK)
    footer(s, 10)


def rules_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "RULES YAML — NGHỊ ĐỊNH 168/2024", size=24, bold=True, color=WHITE)
    add_text(
        s,
        Inches(0.5),
        Inches(1.15),
        Inches(12.3),
        Inches(0.7),
        "Mỗi rule: điều kiện (LHS) · kết luận xử phạt (RHS) · căn cứ pháp lý · (tuỳ chọn) giai_thich_mau.",
        size=14,
        color=DARK,
    )
    card = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.9), Inches(12.3), Inches(4.5))
    fill_shape(card, LIGHT)
    add_text(
        s,
        Inches(0.8),
        Inches(2.1),
        Inches(11.8),
        Inches(4.0),
        [
            "Ví dụ (ô tô · nồng độ cồn mức 3) — khớp source rules/nong_do_con.yaml:",
            "",
            "• id: R_CON_OTO_MUC3   ·   nhom: nong_do_con   ·   ap_dung_loai_xe: [o_to]",
            "• LHS: chiso.nongDoCon_khiTho > 0.4   OR   chiso.nongDoCon_mau > 80",
            "• RHS: tiền 30–40 triệu; tước GPLX 22–24 tháng",
            "• can_cu: NĐ 168/2024/NĐ-CP · Điều 6 · Khoản 11 · điểm a",
            "",
            "Meta-rule (vd. gây tai nạn): chọn mức max trong khung — không nhân hệ số tuỳ ý.",
            "Phạm vi hiện hỗ trợ: cồn, tốc độ, mũ BH, dây AT, đèn đỏ, GPLX + tình tiết tăng/giảm nặng.",
        ],
        size=14,
        color=DARK,
    )
    footer(s, 11)


def digitize_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "QUY TRÌNH SỐ HÓA TRI THỨC", size=26, bold=True, color=WHITE)
    steps = [
        ("1", "Văn bản NĐ", "Nguồn Nghị định\n168/2024"),
        ("2", "Segment &\nExtract", "scripts /\nacquisition"),
        ("3", "YAML Rules", "rules/*.yaml\n+ can_cu"),
        ("4", "Legal-Onto", "concepts.yaml\nkeyphrases"),
        ("5", "Load KB", "Reasoner +\nKnowledgeGraph"),
    ]
    for i, (n, t, d) in enumerate(steps):
        x = Inches(0.4 + i * 2.55)
        c = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(2.35), Inches(3.2))
        fill_shape(c, NAVY if i % 2 == 0 else TEAL)
        add_text(s, x, Inches(2.25), Inches(2.35), Inches(0.5), n, size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.1), Inches(2.9), Inches(2.15), Inches(0.9), t, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.1), Inches(3.9), Inches(2.15), Inches(1.0), d, size=12, color=RGBColor(0xD0, 0xE4, 0xEA), align=PP_ALIGN.CENTER)
    add_text(
        s,
        Inches(0.5),
        Inches(5.6),
        Inches(12.3),
        Inches(0.8),
        "Lưu ý: độ chính xác số hóa luật (~82.6% theo báo cáo) khác với độ chính xác kết luận trên testset (94.7%).",
        size=14,
        color=ACCENT,
    )
    footer(s, 12)


def engine_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "ENGINE SUY DIỄN", size=26, bold=True, color=WHITE)
    add_text(
        s,
        Inches(0.5),
        Inches(1.3),
        Inches(12.3),
        Inches(5.0),
        [
            "Mạng tính toán (M, R): M = biến/thuộc tính bài toán; R = hàm suy (Funcs) + luật điều kiện.",
            "Mô hình bài toán (H, Goal): H = giả thuyết từ facts; Goal = thuộc tính cần suy (vd. mức vượt tốc độ).",
            "",
            "Trong source (Reasoner.infer):",
            "  1) solve_and_apply — A* trên deduction network (suy thuộc tính trung gian)",
            "  2) deduce_rules — forward chaining kích hoạt Rule khớp WM",
            "  3) apply_meta — meta-rule tình tiết → chọn min / max / trung bình khung",
            "  4) aggregate — gộp phạt + Trace KL",
            "",
            "Lời giải tối thiểu: chuỗi suy luận có căn cứ, không dư thừa; mọi kết luận truy xuất được.",
        ],
        size=15,
        color=DARK,
    )
    footer(s, 13)


def cycle_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "CHU TRÌNH SUY DIỄN", size=26, bold=True, color=WHITE)
    steps = [
        ("INPUT", "Câu NL /\nfacts sau clarify"),
        ("WM", "Working\nMemory"),
        ("A*", "Suy thuộc tính\n(Funcs)"),
        ("RULES", "Forward\nchaining"),
        ("META", "Tăng/giảm\nnặng"),
        ("KL", "Kết luận +\nTrace"),
    ]
    for i, (t, d) in enumerate(steps):
        x = Inches(0.35 + i * 2.15)
        c = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, Inches(2.3), Inches(2.0), Inches(2.8))
        fill_shape(c, TEAL if i % 2 == 0 else NAVY)
        add_text(s, x, Inches(2.6), Inches(2.0), Inches(0.6), t, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.1), Inches(3.4), Inches(1.8), Inches(1.3), d, size=13, color=WHITE, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_text(s, x + Inches(1.85), Inches(3.4), Inches(0.35), Inches(0.4), "›", size=22, bold=True, color=ACCENT)
    add_text(
        s,
        Inches(0.5),
        Inches(5.5),
        Inches(12.3),
        Inches(0.9),
        "Clarify loop: thiếu loại xe / khung cồn–tốc độ → hỏi UI → continue_with(facts) — không chạy lại NLU.",
        size=14,
        color=DARK,
    )
    footer(s, 14)


def astar_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "A* VÀ GỘP PHẠT", size=26, bold=True, color=WHITE)
    left = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.3), Inches(6.0), Inches(5.0))
    fill_shape(left, LIGHT)
    add_text(s, Inches(0.8), Inches(1.5), Inches(5.4), Inches(0.4), "Thuật toán A*", size=18, bold=True, color=TEAL)
    add_text(
        s,
        Inches(0.8),
        Inches(2.2),
        Inches(5.4),
        Inches(3.8),
        [
            "f(n) = g(n) + h(n)",
            "• g(n): chi phí đường đi thực tế",
            "• h(n): heuristic ước lượng",
            "",
            "Mục tiêu: chuỗi Func kích hoạt",
            "tối ưu để suy thuộc tính còn thiếu",
            "(vd. mức vượt tốc độ km/h / %).",
            "",
            "Ghi TraceStep kind = ASTAR / FUNC.",
        ],
        size=14,
        color=DARK,
    )
    right = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.3), Inches(6.0), Inches(5.0))
    fill_shape(right, LIGHT)
    add_text(s, Inches(7.1), Inches(1.5), Inches(5.4), Inches(0.4), "Tổng hợp gộp phạt (code)", size=18, bold=True, color=ACCENT)
    add_text(
        s,
        Inches(7.1),
        Inches(2.2),
        Inches(5.4),
        Inches(3.8),
        [
            "• Cộng dồn tiền các rule đã fire",
            "• Hình phạt bổ sung: lấy mức nặng nhất",
            "• Meta tăng nặng → chon_muc = max",
            "• Meta giảm nhẹ → chon_muc = min",
            "• Tăng & giảm cùng lúc → bù trừ (TB khung)",
            "",
            "⚠ Không nhân hệ số phạt tùy ý",
            "  (khác một số mô tả slide cũ).",
            "• Kết quả: tong_tien + can_cu đầy đủ",
        ],
        size=14,
        color=DARK,
    )
    footer(s, 15)


def llm_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "TÍCH HỢP LLM & ĐỒ THỊ CÂU HỎI", size=24, bold=True, color=WHITE)
    add_text(
        s,
        Inches(0.5),
        Inches(1.2),
        Inches(12.3),
        Inches(5.2),
        [
            "Bài toán 1 — Tình tiết gián tiếp:",
            "  LLM / KeywordInferrer chọn trong closed vocabulary + confidence; < 0.6 → cần xác nhận.",
            "",
            "Bài toán 2 — Trích facts:",
            "  LLMExtractor | HeuristicExtractor → JSON facts → grounding schema.",
            "",
            "Legal-Onto question-graph:",
            "  subgraph_match (TF-IDF phrase) → map_to_facts bổ sung slot trống (không ghi đè).",
            "",
            "Providers (sidebar Streamlit):",
            "  • OpenAI (GPT)  ·  Anthropic (Claude)  ·  Local OpenAI-compat (Qwen/Colab/vLLM)",
            "  · Heuristic offline  ·  Tự động theo .env",
            "",
            "Cloud LLM thường chính xác NLU hơn Local 7B; engine phạt vẫn tất định khi facts đúng.",
        ],
        size=14,
        color=DARK,
    )
    footer(s, 16)


def flow_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "LUỒNG DỮ LIỆU TỔNG THỂ", size=26, bold=True, color=WHITE)
    flow = [
        "Người dùng (Streamlit)",
        "↓",
        "NLUPipeline  →  facts + nlu_meta (matched_concepts, inferred, warnings)",
        "↓",
        "needed_clarifications?  →  hỏi radio UI  →  apply_clarification",
        "↓",
        "Reasoner.infer(facts)  →  KetQua + Trace",
        "↓",
        "render_explanation + Nodes graph (Legal-Onto / reasoning subgraph)",
        "↓",
        "Hiển thị: mức phạt · căn cứ · facts · trace · đồ thị",
    ]
    add_text(s, Inches(1.0), Inches(1.3), Inches(11.3), Inches(5.5), flow, size=16, color=DARK, align=PP_ALIGN.CENTER)
    footer(s, 17)


def ui_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "GIAO DIỆN STREAMLIT (CẬP NHẬT)", size=24, bold=True, color=WHITE)
    cards = [
        ("Tab Tư vấn", "Nhập tình huống · Phân tích · Clarify · Giải thích mức phạt · Facts · Trace"),
        ("Tab Nodes graph", "Legal-Onto tương tác (pyvis): EXHIBITS / HAS_KEYPHRASE / MAPS_TO + độ tin"),
        ("Chuỗi suy diễn", "INPUT → concept → fact → WM → RULE → META → Kết luận (hover cạnh)"),
        ("Cấu hình LLM", "Sidebar: provider · API key/.env · model (gpt-4o-mini, Claude, Local)"),
    ]
    for i, (t, b) in enumerate(cards):
        x, y = (0.5 + (i % 2) * 6.4, 1.3 + (i // 2) * 2.6)
        c = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(6.0), Inches(2.3))
        fill_shape(c, LIGHT)
        add_text(s, Inches(x + 0.3), Inches(y + 0.35), Inches(5.4), Inches(0.45), t, size=16, bold=True, color=TEAL)
        add_text(s, Inches(x + 0.3), Inches(y + 1.0), Inches(5.4), Inches(1.0), b, size=14, color=DARK)
    footer(s, 18)


def eval_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "ĐÁNH GIÁ HIỆU NĂNG", size=26, bold=True, color=WHITE)
    add_text(
        s,
        Inches(0.5),
        Inches(1.1),
        Inches(12),
        Inches(0.4),
        "19 ca · HeuristicExtractor (offline) · NĐ 168/2024  —  nguồn: eval/report.py",
        size=13,
        color=MUTED,
    )
    metrics = [
        ("Kết luận đúng tập HV", "94.7%"),
        ("Đúng tổng tiền phạt", "94.7%"),
        ("F1 trích Facts", "99.2%"),
        ("Precision / Recall", "100% / 98.7%"),
        ("NLU trung bình", "0.24 ms"),
        ("Engine trung bình", "0.17 ms"),
    ]
    for i, (k, v) in enumerate(metrics):
        x = Inches(0.5 + (i % 3) * 4.2)
        y = Inches(1.7 + (i // 3) * 2.2)
        c = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, Inches(3.9), Inches(1.9))
        fill_shape(c, WHITE)
        c.line.color.rgb = TEAL
        add_text(s, x + Inches(0.2), y + Inches(0.35), Inches(3.5), Inches(0.4), k, size=13, color=MUTED, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.2), y + Inches(0.85), Inches(3.5), Inches(0.7), v, size=26, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    footer(s, 19)


def conclusion_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "KẾT LUẬN & HẠN CHẾ", size=26, bold=True, color=WHITE)
    add_text(
        s,
        Inches(0.5),
        Inches(1.2),
        Inches(12.3),
        Inches(5.2),
        [
            "✔ Engine ký hiệu tất định: Facts đúng → kết luận & căn cứ luôn đúng. Sai số đến từ NLU.",
            "",
            "✘ Ca c19: “chạy 75 km/h trong khu dân cư” — chưa suy giới hạn 50 km/h từ loại khu vực.",
            "   → bổ sung luật suy giới hạn khu vực hoặc dùng LLMExtractor.",
            "",
            "⚡ Thời gian suy diễn ~ms — phù hợp chatbot thời gian thực.",
            "🔁 Heuristic tái lập 100%; LLM nâng recall câu khẩu ngữ nhưng phụ thuộc API/model.",
            "",
            "⚠ Rủi ro: phủ sóng luật chưa toàn NĐ 168; chi phí token LLM; tiếng Việt đa phương ngữ.",
            "",
            "✅ Hướng phát triển: mở rộng KH + rules; đánh giá head-to-head OpenAI vs Local;",
            "   tối ưu chi phí; làm giàu quan hệ ontology (R_hier / R_con) nếu cần.",
        ],
        size=14,
        color=DARK,
    )
    footer(s, 20)


def refs_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.5), "TÀI LIỆU THAM KHẢO", size=26, bold=True, color=WHITE)
    add_text(
        s,
        Inches(0.5),
        Inches(1.2),
        Inches(12.3),
        Inches(5.5),
        [
            "1. Phạm, V. T., & cộng sự. Legal-Onto: Ontology-based KR for Vietnamese Traffic Law.",
            "2. Phạm, V. T. Neuro-symbolic Reasoning for Legal Expert Systems in Vietnamese Context.",
            "3. Phạm, V. T., & Nguyễn, H. M. Combining LLM and Symbolic Engine for Legal QA. KSE.",
            "4. Nghị định 168/2024/NĐ-CP — xử phạt VPHC lĩnh vực giao thông đường bộ (KB chính).",
            "5. Nghị định 100/2019/NĐ-CP & 123/2021/NĐ-CP — văn bản tiền thân / tham chiếu lịch sử.",
            "6. Nguyen, D. Q., et al. PhoBERT. Findings of EMNLP 2020.",
            "7. underthesea · VnCoreNLP — công cụ NLP tiếng Việt.",
            "8. Source dự án: traffic_es/ (NLU, engine, knowledge, llm, viz) · app/streamlit_app.py",
        ],
        size=14,
        color=DARK,
    )
    footer(s, 21)


def thanks_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, W, H)
    fill_shape(bg, NAVY)
    accent = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, Inches(4.6), W, Inches(0.1))
    fill_shape(accent, TEAL)
    add_text(s, Inches(0.5), Inches(2.5), Inches(12.3), Inches(1.0), "CẢM ƠN THẦY CÔ VÀ CÁC BẠN!", size=36, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(
        s,
        Inches(0.5),
        Inches(5.0),
        Inches(12.3),
        Inches(1.0),
        [
            "CS2307.CH201 · Nhóm 4 · Hệ chuyên gia tư vấn VPHC giao thông dựa trên Reasoning",
            "Demo: streamlit run app/streamlit_app.py",
        ],
        size=14,
        color=RGBColor(0xA8, 0xC5, 0xD4),
        align=PP_ALIGN.CENTER,
    )


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    title_slide(prs)
    members_slide(prs)
    overview_slide(prs)
    goals_slide(prs)
    principles_slide(prs)
    arch_slide(prs)
    legal_onto_slide(prs)
    concept_slide(prs)
    catalog_slide(prs)
    kg_slide(prs)
    rules_slide(prs)
    digitize_slide(prs)
    engine_slide(prs)
    cycle_slide(prs)
    astar_slide(prs)
    llm_slide(prs)
    flow_slide(prs)
    ui_slide(prs)
    eval_slide(prs)
    conclusion_slide(prs)
    refs_slide(prs)
    thanks_slide(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes), slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
